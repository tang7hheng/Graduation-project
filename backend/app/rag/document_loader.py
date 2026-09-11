"""Multi-format document loader dispatcher.

Supported formats: .pdf, .docx, .xlsx (Q/A FAQ or plain rows), .json (FAQ list), .md, .txt
Output: List[langchain_core.documents.Document] with metadata {source, doc_id, page?}
"""
import json
import logging
from pathlib import Path
from typing import Optional

import pandas as pd
from langchain_core.documents import Document

log = logging.getLogger(__name__)


SUPPORTED_EXT = {".pdf", ".docx", ".xlsx", ".json", ".md", ".txt"}


def detect_doc_type(filename: str) -> str:
    ext = Path(filename).suffix.lower()
    return ext.lstrip(".") or "unknown"


def load_file(file_path: str | Path, doc_id: str) -> list[Document]:
    """Dispatch to the correct loader based on file extension.

    Args:
        file_path: absolute path to the source file
        doc_id: unique id used to group chunks in the vector store for later deletion
    """
    p = Path(file_path)
    if not p.exists():
        raise FileNotFoundError(p)

    ext = p.suffix.lower()
    if ext not in SUPPORTED_EXT:
        raise ValueError(f"Unsupported file type: {ext}")

    if ext == ".pdf":
        docs = _load_pdf(p)
    elif ext == ".docx":
        docs = _load_docx(p)
    elif ext == ".xlsx":
        docs = _load_xlsx(p)
    elif ext == ".json":
        docs = _load_json_faq(p)
    elif ext in (".md", ".txt"):
        docs = _load_text(p)
    else:
        docs = []

    # Attach metadata
    for d in docs:
        d.metadata = d.metadata or {}
        d.metadata.setdefault("source", p.name)
        d.metadata.setdefault("doc_id", doc_id)

    log.info("Loaded %d chunk(s) from %s (doc_id=%s)", len(docs), p.name, doc_id)
    return docs


def _load_pdf(p: Path) -> list[Document]:
    from pypdf import PdfReader

    reader = PdfReader(str(p))
    docs = []
    for i, page in enumerate(reader.pages, start=1):
        text = (page.extract_text() or "").strip()
        if text:
            docs.append(Document(page_content=text, metadata={"page": i, "source": p.name}))
    return docs


def _load_docx(p: Path) -> list[Document]:
    import docx

    doc = docx.Document(str(p))
    paragraphs = [para.text.strip() for para in doc.paragraphs if para.text.strip()]
    text = "\n\n".join(paragraphs)
    if not text:
        return []
    # Keep as single doc; chunker will split
    return [Document(page_content=text, metadata={"source": p.name})]


def _load_xlsx(p: Path) -> list[Document]:
    """If first two columns look like Q/A, treat as FAQ; else treat each row as one chunk."""
    try:
        df = pd.read_excel(p, sheet_name=0, engine="openpyxl")
    except Exception as e:
        log.warning("Excel read failed, fallback empty: %s", e)
        return []

    if len(df.columns) >= 2:
        c0, c1 = str(df.columns[0]).lower(), str(df.columns[1]).lower()
        qa_like = c0 in {"question", "q", "问题"} and c1 in {"answer", "a", "答案", "回答"}
    else:
        qa_like = False

    docs = []
    if qa_like:
        for _, row in df.iterrows():
            q = str(row.iloc[0]).strip()
            a = str(row.iloc[1]).strip()
            if not q:
                continue
            content = f"问:{q}\n答:{a}"
            docs.append(Document(page_content=content, metadata={"source": p.name, "is_faq": True}))
    else:
        for _, row in df.iterrows():
            row_text = "\n".join([f"{c}: {row[c]}" for c in df.columns if pd.notna(row[c])])
            if row_text.strip():
                docs.append(Document(page_content=row_text, metadata={"source": p.name}))
    return docs


def _load_json_faq(p: Path) -> list[Document]:
    """JSON FAQ format: list of {"question": "...", "answer": "..."} or {"q":..., "a":...}."""
    with open(p, "r", encoding="utf-8") as f:
        data = json.load(f)

    if not isinstance(data, list):
        raise ValueError("JSON FAQ file must be a list of {question, answer} objects")

    docs = []
    for item in data:
        if not isinstance(item, dict):
            continue
        q = item.get("question") or item.get("q") or ""
        a = item.get("answer") or item.get("a") or ""
        q = str(q).strip()
        a = str(a).strip()
        if not q:
            continue
        docs.append(
            Document(
                page_content=f"问:{q}\n答:{a}",
                metadata={"source": p.name, "is_faq": True},
            )
        )
    return docs


def _load_text(p: Path) -> list[Document]:
    with open(p, "r", encoding="utf-8") as f:
        text = f.read()
    if not text.strip():
        return []
    return [Document(page_content=text, metadata={"source": p.name})]
